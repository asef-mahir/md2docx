"""
Folder Watcher — Auto Study Guide Generator
-------------------------------------------
Watches ./inbox/ for new .md files.
Each new .md file is converted into a styled .docx in ./outbox/.
The original .md is moved to ./done/ afterwards.

Usage:
    pip install python-docx watchdog
    python watcher.py
"""

import re
import time
import shutil
from pathlib import Path
from datetime import datetime

from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ------------- Folder layout -------------
ROOT    = Path(__file__).parent
INBOX   = ROOT / "inbox"
OUTBOX  = ROOT / "outbox"
DONE    = ROOT / "done"
LOG     = ROOT / "conversion_log.txt"

for d in (INBOX, OUTBOX, DONE):
    d.mkdir(exist_ok=True)


# ------------- Style config -------------
ACCENT_COLOR = RGBColor(0x1F, 0x3A, 0x93)
CODE_BG      = "F2F2F2"
QUOTE_COLOR  = RGBColor(0x55, 0x55, 0x55)


def shade(paragraph, hex_color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    pPr.append(shd)


def add_runs_with_formatting(paragraph, text):
    token_re = re.compile(r'(\*\*.+?\*\*|\*.+?\*)')
    for token in token_re.split(text):
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            r = paragraph.add_run(token[2:-2]); r.bold = True
        elif token.startswith('*') and token.endswith('*'):
            r = paragraph.add_run(token[1:-1]); r.italic = True
        else:
            paragraph.add_run(token)


def add_title_page(doc, title, subtitle=None):
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run(title)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = ACCENT_COLOR
    if subtitle:
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = s.add_run(subtitle)
        r.italic = True
        r.font.size = Pt(14)
    doc.add_page_break()


def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '999999')
    pBdr.append(bottom)
    pPr.append(pBdr)


def insert_toc_field(doc):
    toc = doc.add_paragraph()
    toc.add_run('[Right-click here in Word → Update Field to generate TOC]').italic = True
    fldChar  = OxmlElement('w:fldChar'); fldChar.set(qn('w:fldCharType'), 'begin')
    instr    = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar'); fldChar3.set(qn('w:fldCharType'), 'end')
    run = toc.add_run()._r
    run.append(fldChar); run.append(instr); run.append(fldChar2); run.append(fldChar3)


def build_docx(input_text, output_path, title=None, subtitle="Personal Study Material"):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)

    add_title_page(doc, title or "Study Guide", subtitle)
    doc.add_heading('Table of Contents', level=1)
    insert_toc_field(doc)
    doc.add_page_break()

    lines = input_text.splitlines()
    i = 0
    table_buffer = []

    def flush_table():
        nonlocal table_buffer
        if not table_buffer:
            return
        rows = [r for r in table_buffer if not re.match(r'^\s*\|[\s\-:|]+\|\s*$', r)]
        cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
        if cells and cells[0]:
            tbl = doc.add_table(rows=len(cells), cols=len(cells[0]))
            tbl.style = 'Light Grid Accent 1'
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row in enumerate(cells):
                for c_idx, val in enumerate(row):
                    cell = tbl.cell(r_idx, c_idx)
                    cell.text = ''
                    p = cell.paragraphs[0]
                    add_runs_with_formatting(p, val)
                    if r_idx == 0:
                        for run in p.runs:
                            run.bold = True
        table_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('|') and stripped.endswith('|'):
            table_buffer.append(stripped)
            i += 1
            continue
        else:
            flush_table()

        if stripped.startswith('```'):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            shade(p, CODE_BG)
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            continue

        if stripped.startswith('#### '):
            doc.add_heading(stripped[5:], level=3)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=0)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            add_runs_with_formatting(p, stripped[2:])
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            add_runs_with_formatting(p, re.sub(r'^\d+\.\s', '', stripped))
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            run = p.add_run(stripped[2:]); run.italic = True
            run.font.color.rgb = QUOTE_COLOR
        elif stripped == '---':
            add_horizontal_rule(doc)
        elif stripped == '':
            pass
        else:
            p = doc.add_paragraph()
            add_runs_with_formatting(p, stripped)

        i += 1

    flush_table()
    doc.save(output_path)


def title_from_md(text, fallback):
    for line in text.splitlines():
        if line.strip().startswith('# '):
            return line.strip()[2:].strip()
    return fallback


def log(msg):
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with LOG.open("a", encoding="utf-8") as f:
        f.write(f"[{stamp}] {msg}\n")
    print(f"[{stamp}] {msg}")


def convert(md_path: Path):
    try:
        text = md_path.read_text(encoding="utf-8")
        title = title_from_md(text, md_path.stem.replace("_", " ").title())
        out_path = OUTBOX / (md_path.stem + ".docx")
        build_docx(text, out_path, title=title)
        shutil.move(str(md_path), str(DONE / md_path.name))
        log(f"✔ {md_path.name} → {out_path.name}")
    except Exception as e:
        log(f"✘ FAILED {md_path.name}: {e}")


class InboxHandler(FileSystemEventHandler):
    def on_created(self, event):
        if event.is_directory:
            return
        p = Path(event.src_path)
        if p.suffix.lower() == ".md":
            # small delay so file finishes writing
            time.sleep(0.8)
            if p.exists():
                convert(p)


def scan_existing():
    for md in INBOX.glob("*.md"):
        convert(md)


if __name__ == "__main__":
    log(f"Watching {INBOX} …  (Ctrl+C to stop)")
    scan_existing()
    observer = Observer()
    observer.schedule(InboxHandler(), str(INBOX), recursive=False)
    observer.start()
    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        observer.stop()
    observer.join()
    log("Watcher stopped.")